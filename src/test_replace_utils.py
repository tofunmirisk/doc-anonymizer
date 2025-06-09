import os
import tempfile
import csv
from docx import Document
from utils.replace_utils import load_replacements_with_types, replace_words_in_docx

def create_sample_csv(path):
    # Create a sample CSV with type|original|fictitious
    with open(path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f, delimiter='|')
        writer.writerow(['Company', 'Acme Corp', 'Pseudo Corp'])
        writer.writerow(['Name', 'John Doe', 'Jane Roe'])
        writer.writerow(['Company', 'Acme Corp', 'Acme Corp'])  # Should be skipped

def create_sample_docx(path):
    doc = Document()
    doc.add_paragraph("Acme Corp is represented by John Doe.")
    doc.save(path)

def test_load_replacements_with_types_and_replace(tmp_path):
    # Setup
    csv_path = tmp_path / "replacements.csv"
    docx_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"
    create_sample_csv(csv_path)
    create_sample_docx(docx_path)

    # Test loading replacements
    replacements = load_replacements_with_types(str(csv_path))
    assert replacements == {'Acme Corp': 'Pseudo Corp', 'John Doe': 'Jane Roe'}

    # Test replacement in docx
    doc = Document(str(docx_path))
    doc = replace_words_in_docx(doc, replacements)
    doc.save(str(output_path))

    # Check output
    doc_out = Document(str(output_path))
    text = "\n".join([p.text for p in doc_out.paragraphs])
    assert "Pseudo Corp" in text
    assert "Jane Roe" in text
    assert "Acme Corp" not in text or "John Doe" not in text

if __name__ == "__main__":
    import pytest
    pytest.main([__file__])