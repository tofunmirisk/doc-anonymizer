import unittest
from docx import Document
from utils.replace_utils import replace_words_in_docx

class TestReplaceWordsInDocx(unittest.TestCase):
    def setUp(self):
        # Create a new document
        self.doc = Document()
        # Add a paragraph to the body
        self.doc.add_paragraph("Hello John Doe\n, your email is john@example.com\n and your code is 12345.")
        # Add a table to the body
        table = self.doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Contact: 555-1234"
        # Add a header
        section = self.doc.sections[0]
        header = section.header
        header_para = header.add_paragraph("Confidential: John Doe, ID 98765")

    def test_replacement(self):
        replacements = {
            "John Doe": "Jane Smith",
            "john@example.com": "jane@sample.com",
            "12345": "54321",
            "555-1234": "999-8888",
            "98765": "11111"
        }
        replace_words_in_docx(self.doc, replacements)

        # Check body paragraph
        body_text = "\n".join([p.text for p in self.doc.paragraphs])
        self.assertIn("Jane Smith", body_text)
        self.assertIn("jane@sample.com", body_text)
        self.assertIn("54321", body_text)
        self.assertNotIn("John Doe", body_text)
        self.assertNotIn("john@example.com", body_text)
        self.assertNotIn("12345", body_text)

        # Check table
        table_text = self.doc.tables[0].cell(0, 0).text
        self.assertEqual(table_text, "Contact: 999-8888")

        # Check header
        header_text = self.doc.sections[0].header.paragraphs[-1].text
        self.assertIn("Jane Smith", header_text)
        self.assertIn("11111", header_text)
        self.assertNotIn("John Doe", header_text)
        self.assertNotIn("98765", header_text)

if __name__ == "__main__":
    unittest.main()