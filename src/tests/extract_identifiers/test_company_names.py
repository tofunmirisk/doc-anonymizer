import os
import pytest
from docx import Document
from extract_identifiers.company_names import extract_company_names
from extract_identifiers.utils import extract_lines

@pytest.fixture
def sample_docx(tmp_path):
    docx_path = tmp_path / "test_multiline_company.docx"
    doc = Document()
    doc.add_paragraph("INSURED:")
    doc.add_paragraph("Sanctuary Housing")
    doc.add_paragraph("Association Ltd")
    doc.add_paragraph("BROKER:")
    doc.add_paragraph("NovaSecure Risk")
    doc.add_paragraph("Partners LLP")
    doc.save(docx_path)
    return str(docx_path)

def test_multiline_preserved_company_name_extraction(sample_docx):
    doc = Document(sample_docx)
    lines = extract_lines(doc)
    companies = extract_company_names(lines)

    expected_1 = "Sanctuary Housing\nAssociation Ltd"
    expected_2 = "NovaSecure Risk\nPartners LLP"

    assert expected_1 in companies, "❌ Failed to extract multi-line insured"
    assert expected_2 in companies, "❌ Failed to extract multi-line broker"

    print("✅ Multi-line company name preservation test passed.")
    print("Extracted:")
    for c in companies:
        print("-", repr(c))