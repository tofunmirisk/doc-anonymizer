from docx import Document
from dotenv import load_dotenv
load_dotenv()

# todo: extract lines from text box (shapes) and footnotes if present
def extract_lines(doc):
    """
    Extract lines from a DOCX file, treating line breaks within a paragraph as new lines.
    :param doc: A python-docx Document object
    :return: List of lines extracted from paragraphs and tables, with line breaks handled
    """
    def para_text_with_breaks(para):
        # Collect text, inserting \n for each line break within the paragraph
        text = ""
        for run in para.runs:
            text += run.text
            # Insert \n for each <w:br/> in the run
            for br in run._element.findall('.//w:br', run._element.nsmap):
                text += "\n"
        return text

    lines = []
    for para in doc.paragraphs:
        para_text = para_text_with_breaks(para).strip()
        if para_text:
            # Split on \n to treat line breaks as new lines
            lines.extend([line.strip() for line in para_text.split('\n') if line.strip()])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para_text = para_text_with_breaks(para).strip()
                    if para_text:
                        lines.extend([line.strip() for line in para_text.split('\n') if line.strip()])
    return lines


def extract_text(doc):
    texts = []
    for para in doc.paragraphs:
        texts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    texts.append(para.text)
    for section in doc.sections:
        header = section.header
        for para in header.paragraphs:
            texts.append(para.text)
        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        texts.append(para.text)
    return texts


def prepare_prompt(text):
    """
    Prepare the prompt for the LLM to extract sensitive values.
    :param text: The text content from the document
    :return: A formatted prompt string
    """

    return f"""
--- CONTRACT CONTENT START ---
{text}
--- CONTRACT CONTENT END ---
"""

def llm_filter(prompt, system_role=None, model="gpt-4o" ):
    import openai, os
    # openai_api_key = "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"  # Your OpenAI API key

    openai_api_key = os.getenv("OPENAI_API_KEY")
    if not openai_api_key:
        raise ValueError("OPENAI_API_KEY environment variable is not set.")

    client = openai.OpenAI(api_key=openai_api_key)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_role if system_role else "You are a helpful assistant for extracting sensitive values for anonymization."},
            {"role": "user", "content": prompt}
        ],
        max_tokens=2048,
        temperature=0
    )
    import ast
    try:
        sensitive_list = ast.literal_eval(response.choices[0].message.content)
        return sensitive_list
    except Exception:
        return [response.choices[0].message.content.strip()]
    

def create_docx(paragraphs, filename):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(filename)
